"""Import 壹思源维修—标准报价表.xlsx into the local SQLite catalog.

One command from the repo root:

    python3 local-backend/scripts/import-standard-quote.py

The workbook and database paths default to the standard file in the repo root
and local-backend/data/yisiyuan-repair.sqlite. DB_FILE overrides the database
when it is not passed explicitly, matching the backend.
"""

import argparse
import json
import os
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path
from zipfile import BadZipFile

try:
    import openpyxl
except ImportError:
    requirements = Path(__file__).resolve().parent.parent / "requirements.txt"
    sys.stderr.write(
        "导入失败：缺少 Python 依赖 openpyxl。"
        f"请执行：python3 -m pip install -r {requirements}\n"
    )
    sys.exit(1)


SERVICE_SHEETS = ("家电维修", "清洗保养安装", "水电管道厨卫")
MARKER_SHEET = "合作报价说明"
SHEET_HEADERS = {
    "家电维修": ("品类", "服务项目", "服务内容", "维修费用", "包含内容", "不含/另计", "质保"),
    "清洗保养安装": ("品类", "服务项目", "规格/服务内容", "维修费用", "包含内容", "不含/另计", "质保"),
    "水电管道厨卫": ("品类", "服务项目", "服务内容", "维修费用", "包含内容", "不含/另计", "质保"),
    "合作套餐": ("套餐名称", "适用合作商", "服务内容", "维修费用", "适用建议"),
    "特殊费用及规则": ("费用类型", "收费标准", "适用说明"),
}
REQUIRED_TABLES = (
    "app_meta",
    "categories",
    "brands",
    "services",
    "extras",
    "staff_users",
    "customers",
    "sessions",
    "receipts",
    "cities",
    "regions",
)
PROTECTED_TABLES = ("staff_users", "customers", "sessions", "receipts", "cities", "regions")
PROTECTED_LABELS = {
    "staff_users": "staffUsers",
    "customers": "customers",
    "sessions": "sessions",
    "receipts": "receipts",
    "cities": "cities",
    "regions": "regions",
}
COLORS = ("#1888ff", "#5865f2", "#12a594", "#e8952f", "#8a63df", "#df5c55")
SCRIPT_DIR = Path(__file__).resolve().parent
LOCAL_BACKEND_DIR = SCRIPT_DIR.parent
REPO_ROOT = LOCAL_BACKEND_DIR.parent
DEFAULT_WORKBOOK = REPO_ROOT / "壹思源维修—标准报价表.xlsx"
DEFAULT_DATABASE = LOCAL_BACKEND_DIR / "data" / "yisiyuan-repair.sqlite"
SOURCE_NAME = "壹思源维修—标准报价表.xlsx"


class QuoteImportError(Exception):
    pass


def parse_price(value):
    text = str(value or "").strip()
    nums = [int(x) for x in re.findall(r"\d+", text)]
    if not nums:
        return 0, 0, text
    low, high = nums[0], nums[1] if len(nums) > 1 else nums[0]
    return low, high, text


def parse_warranty(value):
    match = re.search(r"\d+", str(value or ""))
    return int(match.group()) if match else 0


def parse_extra_amount(value):
    text = str(value or "").strip()
    price_range = re.search(r"(\d+)\s*-\s*(\d+)\s*元", text)
    if price_range:
        return round((int(price_range.group(1)) + int(price_range.group(2))) / 2)
    single_price = re.search(r"(?:最低加收|加收|^)(\d+)\s*元", text)
    if single_price:
        return int(single_price.group(1))
    starting_price = re.search(r"(\d+)\s*元起", text)
    return int(starting_price.group(1)) if starting_price else 0


def blank(value):
    return value is None or str(value).strip() == ""


def cells(row, width):
    values = list(row[:width]) if row else []
    while len(values) < width:
        values.append(None)
    return values


def text_at(row, index):
    value = row[index] if index < len(row) else None
    return "" if blank(value) else str(value).strip()


def require_text(row, index, sheet_name, row_number, label):
    value = text_at(row, index)
    if not value:
        raise QuoteImportError(f"工作表「{sheet_name}」第 {row_number} 行缺少{label}")
    return value


def header_of(sheet):
    row = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True), None)
    return tuple(text_at(row or (), index) for index in range(len(row or ())))


def assert_standard_workbook(book):
    missing = [name for name in (MARKER_SHEET, *SHEET_HEADERS) if name not in book.sheetnames]
    if missing:
        joined = "、".join(missing)
        raise QuoteImportError(f"不是标准报价表，缺少工作表：{joined}。请使用「{SOURCE_NAME}」，不要导入其他 Excel。")
    marker = next(book[MARKER_SHEET].iter_rows(min_row=1, max_row=1, max_col=1, values_only=True), (None,))
    if "壹思源维修" not in str(marker[0] or ""):
        raise QuoteImportError(f"不是标准报价表：说明页标题应包含「壹思源维修」。请使用「{SOURCE_NAME}」。")
    for sheet_name, expected in SHEET_HEADERS.items():
        actual = header_of(book[sheet_name])
        if actual[: len(expected)] != expected:
            expected_text = "、".join(expected)
            raise QuoteImportError(f"不是标准报价表：工作表「{sheet_name}」第 4 行表头应为「{expected_text}」。")


def rows_from_workbook(path):
    try:
        book = openpyxl.load_workbook(path, data_only=True)
    except (BadZipFile, OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise QuoteImportError(f"无法读取 Excel：{path}。请使用「{SOURCE_NAME}」。") from exc
    try:
        assert_standard_workbook(book)
        categories = []
        services = []
        category_ids = {}

        def category_id(name):
            if name not in category_ids:
                category_ids[name] = f"cat_{len(category_ids) + 1:03d}"
                categories.append((category_ids[name], name, len(categories)))
            return category_ids[name]

        for sheet_name in SERVICE_SHEETS:
            sheet = book[sheet_name]
            for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
                values = cells(row, 7)
                if all(blank(value) for value in values):
                    continue
                category = require_text(values, 0, sheet_name, row_number, "品类")
                name = require_text(values, 1, sheet_name, row_number, "服务项目")
                low, high, price_text = parse_price(values[3])
                services.append({
                    "id": f"svc_{len(services) + 1:03d}",
                    "category_id": category_id(category),
                    "name": name,
                    "description": text_at(values, 2),
                    "low": low,
                    "high": high,
                    "part": 0,
                    "warranty_days": parse_warranty(values[6]),
                    "includes": text_at(values, 4),
                    "excludes": text_at(values, 5),
                    "price_text": price_text,
                    "meta": json.dumps({"sourceSheet": sheet_name, "sourceRow": row_number}, ensure_ascii=False),
                })

        package_sheet = book["合作套餐"]
        package_category = category_id("合作套餐")
        for row_number, row in enumerate(package_sheet.iter_rows(min_row=5, values_only=True), 5):
            values = cells(row, 5)
            if all(blank(value) for value in values):
                continue
            name = require_text(values, 0, "合作套餐", row_number, "套餐名称")
            partner = text_at(values, 1)
            low, high, price_text = parse_price(values[3])
            services.append({
                "id": f"svc_{len(services) + 1:03d}",
                "category_id": package_category,
                "name": name,
                "description": text_at(values, 2),
                "low": low,
                "high": high,
                "part": 0,
                "warranty_days": 0,
                "includes": f"适用合作商：{partner}" if partner else "",
                "excludes": text_at(values, 4),
                "price_text": price_text,
                "meta": json.dumps({"sourceSheet": "合作套餐", "sourceRow": row_number}, ensure_ascii=False),
            })

        extras = []
        sheet = book["特殊费用及规则"]
        for row_number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            values = cells(row, 3)
            if all(blank(value) for value in values):
                continue
            name = require_text(values, 0, "特殊费用及规则", row_number, "费用类型")
            price_text = text_at(values, 1)
            extras.append((
                f"extra_{len(extras) + 1:03d}",
                name,
                parse_extra_amount(price_text),
                price_text,
                text_at(values, 2),
            ))
    finally:
        book.close()

    if not categories or not services or not extras:
        raise QuoteImportError("标准报价表没有可导入的品类、服务项目或特殊费用，已停止，数据库未修改。")
    return categories, services, extras


def connect(database):
    try:
        db = sqlite3.connect(database, timeout=30)
    except sqlite3.OperationalError as exc:
        raise QuoteImportError(f"无法打开数据库 {database}：{exc}") from exc
    db.execute("PRAGMA busy_timeout=30000")
    return db


def assert_database_ready(database):
    if not database.exists():
        raise QuoteImportError(
            f"数据库不存在：{database}。请先启动一次后端（在 local-backend 目录执行 node src/server.mjs）生成数据库，或用参数指定已有的 SQLite 文件。"
        )
    if not database.is_file():
        raise QuoteImportError(f"数据库路径不是文件：{database}")
    db = connect(database)
    try:
        names = {row[0] for row in db.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    finally:
        db.close()
    missing = [name for name in REQUIRED_TABLES if name not in names]
    if missing:
        joined = "、".join(missing)
        raise QuoteImportError(f"数据库还没有完成初始化，缺少表：{joined}。请先启动一次后端，再重新导入。")


def backup_database(database, moment):
    stamp = moment.strftime("%Y%m%d-%H%M%S")
    backup = database.with_name(f"{database.stem}.before-excel-{stamp}{database.suffix}")
    index = 2
    while backup.exists():
        backup = database.with_name(f"{database.stem}.before-excel-{stamp}-{index}{database.suffix}")
        index += 1
    source = connect(database)
    try:
        dest = sqlite3.connect(backup)
        try:
            source.backup(dest)
        finally:
            dest.close()
    except sqlite3.OperationalError as exc:
        backup.unlink(missing_ok=True)
        if "locked" in str(exc).lower():
            raise QuoteImportError("数据库正被后端占用，备份没有完成，数据未修改。") from exc
        raise QuoteImportError(f"备份失败，数据未修改。{exc}") from exc
    except Exception:
        backup.unlink(missing_ok=True)
        raise
    finally:
        source.close()
    if not backup.is_file() or backup.stat().st_size <= 0:
        backup.unlink(missing_ok=True)
        raise QuoteImportError(f"备份失败，没有留下可用文件：{backup}")
    return backup


def column_names(db, table):
    return {row[1] for row in db.execute(f"PRAGMA table_info({table})")}


def ensure_columns(db):
    if "price_text" not in column_names(db, "services"):
        db.execute("ALTER TABLE services ADD COLUMN price_text TEXT")
    if "price" not in column_names(db, "extras"):
        db.execute("ALTER TABLE extras ADD COLUMN price TEXT")
    if "scene" not in column_names(db, "extras"):
        db.execute("ALTER TABLE extras ADD COLUMN scene TEXT")


def protected_counts(db):
    return {PROTECTED_LABELS[table]: db.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0] for table in PROTECTED_TABLES}


def apply_import(database, parsed, source, imported_at):
    categories, services, extras = parsed
    db = connect(database)
    try:
        db.execute("PRAGMA foreign_keys=ON")
        ensure_columns(db)
        before = protected_counts(db)
        try:
            db.execute("BEGIN IMMEDIATE")
            db.execute("DELETE FROM services")
            db.execute("DELETE FROM extras")
            db.execute("DELETE FROM brands")
            db.execute("DELETE FROM categories")
            for category_id, name, sort_order in categories:
                tone = COLORS[sort_order % len(COLORS)]
                db.execute(
                    "INSERT INTO categories (id,name,enabled,sort_order,tone,accent,image,meta) VALUES (?,?,?,?,?,?,?,?)",
                    (
                        category_id,
                        name,
                        1,
                        sort_order,
                        tone,
                        "#f5a623",
                        "",
                        json.dumps({"source": SOURCE_NAME}, ensure_ascii=False),
                    ),
                )
            all_category_ids = [item[0] for item in categories]
            db.execute(
                "INSERT INTO brands (id,name,latin,color,factor,category_ids,categories,special,logo,meta) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (
                    "brand_general",
                    "通用品牌",
                    "GENERAL",
                    "#14766f",
                    1,
                    json.dumps(all_category_ids),
                    json.dumps([item[1] for item in categories], ensure_ascii=False),
                    0,
                    "",
                    "{}",
                ),
            )
            for service in services:
                db.execute(
                    "INSERT INTO services (id,category_id,name,description,low,high,part,warranty_days,includes,excludes,price_text,enabled,meta) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        service["id"],
                        service["category_id"],
                        service["name"],
                        service["description"],
                        service["low"],
                        service["high"],
                        service["part"],
                        service["warranty_days"],
                        service["includes"],
                        service["excludes"],
                        service["price_text"],
                        1,
                        service["meta"],
                    ),
                )
            db.executemany("INSERT INTO extras (id,name,amount,enabled,price,scene) VALUES (?,?,?,1,?,?)", extras)
            counts = {
                "categories": len(categories),
                "services": len(services),
                "extras": len(extras),
                "brands": 1,
            }
            db.execute("INSERT OR REPLACE INTO app_meta (key,value) VALUES ('quote_import_source',?)", (source,))
            db.execute("INSERT OR REPLACE INTO app_meta (key,value) VALUES ('quote_imported_at',?)", (imported_at,))
            db.execute(
                "INSERT OR REPLACE INTO app_meta (key,value) VALUES ('quote_import_counts',?)",
                (json.dumps(counts, ensure_ascii=False),),
            )
            seeded = db.execute("SELECT value FROM app_meta WHERE key = 'seeded'").fetchone()
            if seeded is None:
                db.execute("INSERT INTO app_meta (key,value) VALUES ('seeded',?)", (imported_at,))
            after = protected_counts(db)
            if before != after:
                raise QuoteImportError("导入中止：人员、收据、客户、会话、城市或区域数量发生变化，已回滚。")
            db.commit()
        except Exception:
            db.rollback()
            raise
        return counts, after
    except sqlite3.OperationalError as exc:
        message = str(exc).lower()
        if "locked" in message:
            raise QuoteImportError("数据库正被后端占用。请等当前保存完成后再导入。") from exc
        raise QuoteImportError(f"写入数据库失败，报价目录未修改。{exc}") from exc
    finally:
        db.close()


def default_database():
    configured = os.environ.get("DB_FILE", "").strip()
    if configured:
        return Path(configured).expanduser().resolve()
    return DEFAULT_DATABASE


def parse_args(argv):
    parser = argparse.ArgumentParser(description="把标准报价表导入本地 SQLite。只重建品类、服务项目、特殊费用和通用品牌。")
    parser.add_argument("workbook", nargs="?", type=Path, help=f"标准报价表，默认 {DEFAULT_WORKBOOK}")
    parser.add_argument("database", nargs="?", type=Path, help="SQLite 文件。省略时使用环境变量 DB_FILE，否则用后端默认库。")
    args = parser.parse_args(argv)
    workbook = (args.workbook or DEFAULT_WORKBOOK).expanduser().resolve()
    if args.database:
        database = args.database.expanduser().resolve()
    else:
        database = default_database()
    return workbook, database


def run(argv=None):
    workbook, database = parse_args(argv)
    if not workbook.is_file():
        raise QuoteImportError(f"找不到标准报价表：{workbook}")
    assert_database_ready(database)
    parsed = rows_from_workbook(workbook)
    moment = datetime.now(timezone.utc).replace(microsecond=0)
    imported_at = moment.strftime("%Y-%m-%dT%H:%M:%SZ")
    backup = backup_database(database, moment)
    try:
        counts, preserved = apply_import(database, parsed, str(workbook), imported_at)
    except Exception:
        sys.stderr.write(f"导入未完成，原数据库未修改。备份仍保留在：{backup}\n")
        raise
    return {
        "importedAt": imported_at,
        "source": str(workbook),
        "backup": str(backup),
        **counts,
        "preserved": preserved,
    }


def print_success(result):
    preserved = result["preserved"]
    print("标准报价导入成功")
    print(f"时间: {result['importedAt']}")
    print(f"来源: {result['source']}")
    print(f"备份: {result['backup']}")
    print(f"品类: {result['categories']}")
    print(f"服务项目: {result['services']}")
    print(f"特殊费用: {result['extras']}")
    print(f"品牌: {result['brands']}")
    print(
        "未改动人员 {staffUsers}、收据 {receipts}、客户 {customers}、会话 {sessions}、城市 {cities}、区域 {regions}".format(
            **preserved
        )
    )
    print(json.dumps({"ok": True, **result}, ensure_ascii=False))


def main(argv=None):
    try:
        result = run(argv)
    except QuoteImportError as exc:
        sys.stderr.write(f"导入失败：{exc}\n")
        return 1
    print_success(result)
    return 0


if __name__ == "__main__":
    sys.exit(main())
